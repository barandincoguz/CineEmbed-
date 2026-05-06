"""One builder function per slide. Each builder accepts a `prs` (Presentation)
and adds a fully-built slide. Phase 4 ships stubs; Phase 5 fills them in.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from . import theme, components as C

TOTAL = 12


def _new_blank_slide(prs):
    """Add a blank-layout slide (layout index 6 in default master)."""
    blank_layout = prs.slide_layouts[6]
    return prs.slides.add_slide(blank_layout)


def _content_slide(prs, idx: int, title: str, subtitle: str = ""):
    """Standard content slide: top bar, title, optional subtitle, footer."""
    s = _new_blank_slide(prs)
    C.add_top_bar(s)
    C.add_title(s, title)
    if subtitle:
        C.add_subtitle(s, subtitle)
    C.add_footer(s, idx, TOTAL)
    return s


def slide_01_title(prs):
    s = _new_blank_slide(prs)
    # Orange ribbon at top of slide (full width, ~2 inch tall)
    accent = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(0), Emu(0), theme.SLIDE_W, Inches(2.0),
    )
    accent.line.fill.background()
    accent.fill.solid()
    accent.fill.fore_color.rgb = theme.Colors.PRIMARY

    # Title — large, white-on-orange, in the ribbon
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.55), Inches(12.333), Inches(1.0))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "CineEmbed"
    run.font.name = theme.Fonts.HEADING
    run.font.size = Pt(60)
    run.font.color.rgb = theme.Colors.WHITE

    # Subtitle inside the ribbon
    sub = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12.333), Inches(0.5))
    tfs = sub.text_frame
    tfs.margin_left = tfs.margin_right = Emu(0)
    ps = tfs.paragraphs[0]
    runs = ps.add_run()
    runs.text = "Multi-Modal Unsupervised Embedding of 329,044 Films"
    runs.font.name = theme.Fonts.BODY
    runs.font.size = Pt(20)
    runs.font.color.rgb = theme.Colors.WHITE

    # Document label below ribbon
    label = s.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(12.333), Inches(0.5))
    tfl = label.text_frame
    tfl.margin_left = tfl.margin_right = Emu(0)
    pl = tfl.paragraphs[0]
    runl = pl.add_run()
    runl.text = "Intermediate Progress Report  ·  v1.0"
    runl.font.name = theme.Fonts.BODY
    runl.font.size = Pt(20)
    runl.font.color.rgb = theme.Colors.SECONDARY

    # Authors
    auth = s.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(0.6))
    tfa = auth.text_frame
    tfa.margin_left = tfa.margin_right = Emu(0)
    pa = tfa.paragraphs[0]
    runa = pa.add_run()
    runa.text = "Ahmet Baran Dincoguz   ·   Arda Arvas   ·   Bertan Kaan Kaya"
    runa.font.name = theme.Fonts.BODY
    runa.font.size = Pt(22)
    runa.font.color.rgb = theme.Colors.NEAR_BLACK

    # Course / institution
    course = s.shapes.add_textbox(Inches(0.5), Inches(4.95), Inches(12.333), Inches(0.5))
    tfc = course.text_frame
    pc = tfc.paragraphs[0]
    runc = pc.add_run()
    runc.text = "SENG 474 — Deep Learning · Spring 2026 · TED University · 2026-05-06"
    runc.font.name = theme.Fonts.BODY
    runc.font.size = Pt(16)
    runc.font.color.rgb = theme.Colors.SLATE

    # Repo URL
    repo = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(12.333), Inches(0.4))
    tfr = repo.text_frame
    pr = tfr.paragraphs[0]
    runr = pr.add_run()
    runr.text = "github.com/barandincoguz/CineEmbed-"
    runr.font.name = theme.Fonts.BODY
    runr.font.size = Pt(14)
    runr.font.color.rgb = theme.Colors.SECONDARY


def slide_02_status(prs):
    s = _content_slide(prs, 2, "Project Status: ON TRACK",
                       "Modeling MVP delivered. Hypotheses H1, H2, H3 all PASS.")
    C.add_status_pill(s, "complete", left=Inches(0.5), top=Inches(1.55),
                      width=Inches(2.0), height=Inches(0.45))
    C.add_bullets(s,
        items=[
            "Six runs trained, evaluated against 3 orthogonal label axes (genre · decade · language).",
            "Best deep model dec_z64_k21 reaches genre_NMI = 0.332, +205 % over best non-deep baseline.",
            "Bonus finding: films with missing release date form a coherent latent sub-manifold.",
            "Final-report scope (VAE family, k-sweep, F1/F2 ablations) defined and on schedule.",
        ],
        left=Inches(0.5), top=Inches(2.25), width=Inches(7.0), height=Inches(3.5),
    )
    C.add_headline_number(s, big_text="+205 %",
        caption_text="DEC genre_NMI vs KMeans on raw features (0.332 vs 0.109)",
        left=Inches(7.5), top=Inches(2.20), width=Inches(5.4),
    )


def slide_03_goal(prs):
    s = _content_slide(prs, 3, "Project Goal & Scope",
                       "What we set out to do, on what data.")
    C.add_bullets(s,
        items=[
            "Train unsupervised 64-dim representations of movie metadata.",
            "329,044 films from TMDB + awards records + Wikipedia director bios.",
            "564-dim feature matrix organized into 7 modality blocks.",
            "Recover 3 orthogonal label axes without labels: genre, decade, language.",
            "Evaluate via KMeans clustering of the latent → NMI / ARI.",
        ],
        left=Inches(0.5), top=Inches(1.7), width=Inches(7.5), height=Inches(5.0),
    )
    C.add_table(s,
        header=["Block", "Dim", "Notes"],
        rows=[
            ["numerical", "6",   "popularity, runtime, votes"],
            ["genre",     "22",  "21-way one-hot + flag"],
            ["language",  "31",  "top-30 langs (~99% zero)"],
            ["decade",    "2",   "decade_norm + flag"],
            ["awards",    "6",   "Oscar/BAFTA/Cannes counts"],
            ["text",      "384", "MiniLM-L6-v2 embeddings"],
            ["director",  "113", "bio_pca_64 + lang/country"],
        ],
        left=Inches(8.2), top=Inches(1.65), width=Inches(4.7), height=Inches(3.6),
        col_aligns=["l", "r", "l"],
    )
def slide_04_schedule(prs):    _content_slide(prs, 4, "Slide 4 — schedule (stub)")
def slide_05_data(prs):        _content_slide(prs, 5, "Slide 5 — data (stub)")
def slide_06_arch(prs):        _content_slide(prs, 6, "Slide 6 — architecture (stub)")
def slide_07_mvp_table(prs):   _content_slide(prs, 7, "Slide 7 — MVP table (stub)")
def slide_08_headline(prs):    _content_slide(prs, 8, "Slide 8 — headline (stub)")
def slide_09_topology(prs):    _content_slide(prs, 9, "Slide 9 — topology (stub)")
def slide_10_bonus(prs):       _content_slide(prs, 10, "Slide 10 — bonus (stub)")
def slide_11_plan(prs):        _content_slide(prs, 11, "Slide 11 — plan (stub)")
def slide_12_close(prs):       _content_slide(prs, 12, "Slide 12 — close (stub)")


BUILDERS = [
    slide_01_title, slide_02_status, slide_03_goal, slide_04_schedule,
    slide_05_data, slide_06_arch, slide_07_mvp_table, slide_08_headline,
    slide_09_topology, slide_10_bonus, slide_11_plan, slide_12_close,
]
